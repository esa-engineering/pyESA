# -*- coding: utf-8 -*-
"""Genera FamilyNaming_data.py dai tre workbook di classificazione.

Non fa parte del tool: si lancia a mano (py gen_data.py) quando i BIM manager
modificano gli .xlsx, e riscrive per intero il modulo dati.
"""
import io
import os
import sys
import openpyxl
from openpyxl.utils import range_boundaries

HERE = os.path.dirname(os.path.abspath(__file__))

BOOKS = [
    ('LOADABLE',   'Naming classification-Loadable Families.xlsx'),
    ('SYSTEM',     'Naming classification-System Families.xlsx'),
    ('STRUCTURAL', 'Naming classification-Structural Families.xlsx'),
]


def cell_text(ws, row, col):
    v = ws.cell(row=row, column=col).value
    if v is None:
        return u''
    if isinstance(v, float) and v == int(v):
        v = int(v)
    return unicode(v).strip() if str is bytes else str(v).strip()


def read_table(ws, ref):
    """Righe di una Excel Table, header escluso, come liste di stringhe."""
    c1, r1, c2, r2 = range_boundaries(ref)
    rows = []
    for r in range(r1 + 1, r2 + 1):
        vals = [cell_text(ws, r, c) for c in range(c1, c2 + 1)]
        if any(vals):
            rows.append(vals)
    return rows


def read_notes(ws):
    """Le note in coda a un foglio generatore: colonna A, gruppi separati da vuoti.

    Un gruppo di una riga sola e' un titolo di sezione o un paragrafo libero,
    un gruppo di piu' righe e' intestazione + corpo.
    """
    start = None
    for r in range(1, ws.max_row + 1):
        if cell_text(ws, r, 1).startswith(u'NOTE'):
            start = r
            break
    if start is None:
        return []

    groups = []
    current = []
    for r in range(start, ws.max_row + 1):
        txt = cell_text(ws, r, 1)
        if txt:
            current.append(txt)
        elif current:
            groups.append(current)
            current = []
    if current:
        groups.append(current)

    out = []
    for g in groups:
        if len(g) == 1:
            head = g[0]
            upper = head.replace(u'—', u'').replace(u'-', u'').strip()
            kind = 'section' if upper == upper.upper() else 'para'
            out.append((kind, head if kind == 'section' else u'', head if kind == 'para' else u''))
        else:
            out.append(('item', g[0], u' '.join(g[1:])))
    # l'ultima riga di ogni foglio e' il disclaimer sulle colonne grigie: inutile
    out = [o for o in out if u'colonne grigie' not in o[2].lower()]
    return out


def py_repr(value):
    """Letterale unicode valido sia in IronPython 2.7 sia in CPython 3."""
    if isinstance(value, (list, tuple)):
        return u'(' + u', '.join(py_repr(v) for v in value) + (u',)' if len(value) == 1 else u')')
    bs = chr(92)
    s = value.replace(bs, bs + bs).replace(chr(34), bs + chr(34))
    s = s.replace(chr(10), u" ").replace(chr(13), u"")
    return u'u"' + s + u'"'


def main():
    out = []
    w = out.append

    w(u'# -*- coding: utf-8 -*-')
    w(u'"""Tabelle di classificazione ESA per la nomenclatura di famiglie e tipi.')
    w(u'')
    w(u'GENERATO AUTOMATICAMENTE - non modificare a mano.')
    w(u'Sorgente: i tre file "Naming classification-*.xlsx" accanto a questo modulo.')
    w(u'Per rigenerarlo dopo una modifica agli Excel si lancia lo script gen_data.py')
    w(u'descritto nel README del bundle.')
    w(u'')
    w(u'TABLES  nome tabella Excel -> tupla di righe, ogni riga tupla di stringhe.')
    w(u'        La prima colonna e\' sempre il codice, la seconda l\'etichetta inglese,')
    w(u'        la terza la descrizione italiana, la quarta (dove c\'e\') il TM Code.')
    w(u'NOTES   id foglio -> tupla di (tipo, intestazione, corpo).')
    w(u'PATTERNS  id foglio -> tupla di righe che descrivono il pattern del nome.')
    w(u'"""')
    w(u'')

    tables = {}
    notes = {}
    patterns = {}

    for book_key, fname in BOOKS:
        path = os.path.join(HERE, fname)
        wb = openpyxl.load_workbook(path, data_only=True)

        dv = wb['DV']
        for tname, tobj in dv.tables.items():
            key = u'{0}:{1}'.format(book_key, tname)
            tables[key] = read_table(dv, tobj if isinstance(tobj, str) else tobj.ref)

        for ws in wb.worksheets:
            title = ws.title
            if title in ('Istruzioni', 'DV', '_src'):
                continue
            sheet_id = u'{0}:{1}'.format(book_key, title.split(' ')[0])
            notes[sheet_id] = read_notes(ws)
            pat = []
            for r in (3, 4):
                t = cell_text(ws, r, 1)
                if t and (u'_' in t or t.startswith(u'NOME')):
                    pat.append(t)
            patterns[sheet_id] = pat

    w(u'TABLES = {')
    for key in sorted(tables):
        rows = tables[key]
        w(u'    {0}: ('.format(py_repr(key)))
        for row in rows:
            while row and not row[-1]:
                row = row[:-1]
            w(u'        {0},'.format(py_repr(tuple(row))))
        w(u'    ),')
    w(u'}')
    w(u'')

    w(u'PATTERNS = {')
    for key in sorted(patterns):
        if patterns[key]:
            w(u'    {0}: {1},'.format(py_repr(key), py_repr(tuple(patterns[key]))))
    w(u'}')
    w(u'')

    w(u'NOTES = {')
    for key in sorted(notes):
        w(u'    {0}: ('.format(py_repr(key)))
        for kind, head, body in notes[key]:
            w(u'        ({0}, {1}, {2}),'.format(py_repr(kind), py_repr(head), py_repr(body)))
        w(u'    ),')
    w(u'}')
    w(u'')

    target = os.path.join(HERE, 'FamilyNaming_data.py')
    if len(sys.argv) > 1:
        target = sys.argv[1]
    f = io.open(target, 'w', encoding='utf-8')
    f.write(u'\n'.join(out))
    f.write(u'\n')
    f.close()
    print('written {0}  ({1} tables, {2} sheets)'.format(target, len(tables), len(notes)))


main()
